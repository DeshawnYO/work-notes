"""Local, allowlisted HTML, embedded images, and editable DOCX conversion.

No external image fetching: screenshots live in the database and DOCX itself.
"""
from __future__ import annotations
import base64
import html
import io
import re
import zipfile
from html.parser import HTMLParser
from urllib.parse import urlsplit

TAGS = {'p','div','br','h1','h2','h3','h4','h5','h6','strong','b','em','i','u','s','del',
        'ul','ol','li','blockquote','a','table','thead','tbody','tr','td','th','img','span','pre','code','hr'}
VOID = {'br','img','hr'}
MAX_HTML = 20 * 1024 * 1024


def image_bytes(src):
    match = re.fullmatch(r'data:image/(png|jpeg|jpg|gif|webp);base64,([A-Za-z0-9+/=\s]+)', str(src))
    if not match:
        raise ValueError('图片必须上传或粘贴到本机，不能使用临时地址或远程图片链接')
    try:
        data = base64.b64decode(match[2], validate=False)
        if len(data) > 8 * 1024 * 1024:
            raise ValueError('单张图片不能超过 8 MB')
        from PIL import Image
        with Image.open(io.BytesIO(data)) as img:
            if img.width * img.height > 40_000_000:
                raise ValueError('图片像素过大，请缩小后插入')
            img.verify()
        return data
    except (OSError, base64.binascii.Error) as exc:
        raise ValueError('图片内容损坏，请重新插入') from exc


class Cleaner(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out = []; self.stack = []; self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ('script','style','iframe','object','svg','math'):
            self.skip += 1; return
        if self.skip or tag not in TAGS:
            return
        values = dict(attrs); safe = {}
        if tag == 'a':
            href = values.get('href','')
            if urlsplit(href).scheme.lower() in ('http','https','mailto') or href.startswith('#'):
                safe['href'] = href; safe['rel'] = 'noopener noreferrer'
        if tag == 'img':
            src = values.get('src',''); image_bytes(src)
            safe['src'] = src; safe['alt'] = values.get('alt','图片')
            for key in ('width','height'):
                raw = values.get(key,'')
                if str(raw).isdigit(): safe[key] = str(max(16,min(2400,int(raw))))
        if tag in ('td','th'):
            for key in ('colspan','rowspan'):
                raw = values.get(key,'')
                if str(raw).isdigit(): safe[key] = str(max(1,min(30,int(raw))))
        # Browser commands use inline spans. Keep only basic, portable formatting.
        style = values.get('style','')
        kept = []
        for name,value in re.findall(r'([\w-]+)\s*:\s*([^;]+)',style):
            name=name.lower(); value=value.strip().lower()
            if (name=='font-weight' and value in ('bold','700','600')) or (name=='font-style' and value=='italic') or (name=='text-decoration' and value in ('underline','line-through')) or (name=='text-align' and value in ('left','center','right')):
                kept.append(name+':'+value)
        if kept: safe['style'] = ';'.join(kept)
        self.out.append('<'+tag+''.join(' '+k+'="'+html.escape(v,quote=True)+'"' for k,v in safe.items())+'>')
        if tag not in VOID: self.stack.append(tag)

    def handle_endtag(self,tag):
        if tag in ('script','style','iframe','object','svg','math'):
            self.skip = max(0,self.skip-1); return
        if not self.skip and tag in self.stack:
            while self.stack:
                item=self.stack.pop(); self.out.append('</'+item+'>')
                if item==tag: break

    def handle_data(self,data):
        if not self.skip: self.out.append(html.escape(data))


def clean(value):
    value=str(value or '')
    if len(value.encode('utf-8')) > MAX_HTML: raise ValueError('图文内容超过 20 MB，请减少图片大小')
    p=Cleaner(); p.feed(value); p.close()
    return ''.join(p.out)+''.join('</'+tag+'>' for tag in reversed(p.stack))


class Node:
    def __init__(self,tag='',attrs=None): self.tag=tag; self.attrs=attrs or {}; self.children=[]


class Tree(HTMLParser):
    def __init__(self,value):
        super().__init__(convert_charrefs=True); self.root=Node(); self.stack=[self.root]; self.feed(value)
    def handle_starttag(self,tag,attrs):
        n=Node(tag,dict(attrs)); self.stack[-1].children.append(n)
        if tag not in VOID: self.stack.append(n)
    def handle_endtag(self,tag):
        for i in range(len(self.stack)-1,0,-1):
            if self.stack[i].tag==tag: self.stack=self.stack[:i]; break
    def handle_data(self,data): self.stack[-1].children.append(data)


def plain(value):
    def walk(n):
        if isinstance(n,str): return n
        if n.tag=='img': return '[图片]'
        return ''.join(walk(c) for c in n.children)+ ('\n' if n.tag in ('p','div','br','li','tr','h1','h2','h3','blockquote') else '')
    return walk(Tree(value).root).strip()


def to_docx(container,value):
    from docx.shared import Inches
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from PIL import Image

    def inline(p,node,marks=None):
        marks=dict(marks or {})
        if isinstance(node,str):
            r=p.add_run(node)
            for key,val in marks.items(): setattr(r,key,val)
            return
        tag=node.tag; style=node.attrs.get('style','')
        if tag in ('b','strong') or 'font-weight' in style: marks['bold']=True
        if tag in ('i','em') or 'font-style:italic' in style: marks['italic']=True
        if tag=='u' or 'underline' in style: marks['underline']=True
        if tag=='br': p.add_run().add_break(); return
        if tag=='img':
            data=image_bytes(node.attrs['src'])
            with Image.open(io.BytesIO(data)) as img:
                width=img.width
                if img.format not in ('PNG','JPEG','GIF','BMP','TIFF'):
                    buf=io.BytesIO(); img.convert('RGB').save(buf,format='PNG'); data=buf.getvalue()
            requested=int(node.attrs.get('width',width))
            p.add_run().add_picture(io.BytesIO(data),width=Inches(min(6.25,requested/96)))
            return
        if tag=='a' and node.attrs.get('href'):
            link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),p.part.relate_to(node.attrs['href'],RT.HYPERLINK,is_external=True))
            temp=p.add_run()
            for c in node.children: inline(p,c,marks)
            idx=list(p._p).index(temp._r)
            for item in list(p._p)[idx+1:]:
                p._p.remove(item); link.append(item)
            p._p.remove(temp._r); p._p.append(link); return
        for c in node.children: inline(p,c,marks)

    blocks={'p','div','h1','h2','h3','h4','h5','h6','blockquote','pre'}
    def children(target,node):
        pending=[]
        def flush():
            if pending:
                p=target.add_paragraph()
                for c in pending: inline(p,c)
                pending.clear()
        for child in node.children:
            if isinstance(child,str):
                if child.strip(): pending.append(child)
                continue
            tag=child.tag
            if tag in blocks:
                flush()
                if tag=='div' and any(isinstance(c,Node) and c.tag in blocks|{'table','ul','ol'} for c in child.children): children(target,child); continue
                p=target.add_paragraph(style='Heading '+str(min(4,int(tag[1]))) if re.fullmatch('h[1-6]',tag) else 'Normal')
                if tag=='blockquote': p.paragraph_format.left_indent=Inches(.3)
                for c in child.children: inline(p,c)
            elif tag in ('ul','ol'):
                flush(); add_list(target,child,0)
            elif tag=='table':
                flush(); add_table(target,child)
            elif tag=='hr':
                flush(); target.add_paragraph('────────────────')
            else: pending.append(child)
        flush()

    def add_list(target,node,depth):
        for li in node.children:
            if not isinstance(li,Node) or li.tag!='li': continue
            p=target.add_paragraph(style='List Number' if node.tag=='ol' else 'List Bullet')
            p.paragraph_format.left_indent=Inches(.25+.2*depth)
            for c in li.children:
                if isinstance(c,Node) and c.tag in ('ul','ol'): add_list(target,c,depth+1)
                else: inline(p,c)

    def add_table(target,node):
        def rows(n):
            result=[]
            for c in n.children:
                if isinstance(c,Node):
                    if c.tag=='tr': result.append(c)
                    elif c.tag in ('thead','tbody'): result.extend(rows(c))
            return result
        table_rows=rows(node)
        if not table_rows: return
        occupied={}; placements=[]; maxcol=0
        for ri,row in enumerate(table_rows):
            col=0
            for cell in (c for c in row.children if isinstance(c,Node) and c.tag in ('td','th')):
                while (ri,col) in occupied: col+=1
                rs=int(cell.attrs.get('rowspan',1)); cs=int(cell.attrs.get('colspan',1))
                placements.append((ri,col,rs,cs,cell))
                for r in range(ri,ri+rs):
                    for c in range(col,col+cs): occupied[r,c]=True
                col+=cs; maxcol=max(maxcol,col)
        if maxcol>30 or len(table_rows)>1000: raise ValueError('表格过大，请拆分后保存')
        total=max((r+rs for r,c,rs,cs,cell in placements),default=len(table_rows))
        table=target.add_table(rows=total,cols=maxcol); table.style='Table Grid'; table.autofit=True
        for r,c,rs,cs,node in placements:
            cell=table.cell(r,c)
            if rs>1 or cs>1: cell=cell.merge(table.cell(r+rs-1,c+cs-1))
            children(cell,node)
            while len(cell.paragraphs)>1 and not cell.paragraphs[0].text and not cell.paragraphs[0]._p.xpath('.//w:drawing'):
                first=cell.paragraphs[0]._p; first.getparent().remove(first)
            if node.tag=='th':
                for p in cell.paragraphs:
                    for run in p.runs: run.bold=True
    children(container,Tree(clean(value)).root)


def import_docx(data):
    from docx import Document
    from docx.oxml.ns import qn
    if len(data)>16*1024*1024: raise ValueError('导入文件不能超过 16 MB')
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist())>64*1024*1024: raise ValueError('Word 解压内容过大')
    doc=Document(io.BytesIO(data))
    def run_html(r):
        parts=[]
        for elem in r.iter():
            if elem.tag==qn('w:t'): parts.append(html.escape(elem.text or ''))
            elif elem.tag==qn('w:br'): parts.append('<br>')
            elif elem.tag==qn('w:tab'): parts.append('　')
            elif elem.tag.endswith('}blip'):
                rid=elem.get(qn('r:embed')); part=doc.part.related_parts.get(rid)
                if part:
                    parts.append('<img src="data:'+part.content_type+';base64,'+base64.b64encode(part.blob).decode()+'">')
        out=''.join(parts)
        for prop,tag in [('b','strong'),('i','em'),('u','u')]:
            if r.find('.//'+qn('w:'+prop)) is not None: out='<'+tag+'>'+out+'</'+tag+'>'
        return out
    def blocks(parent):
        output=[]
        for elem in parent:
            if elem.tag==qn('w:p'):
                style=elem.find('./'+qn('w:pPr')+'/'+qn('w:pStyle'))
                st=style.get(qn('w:val'),'') if style is not None else ''
                tag='h'+st[-1] if re.fullmatch(r'Heading[1-6]',st) else 'p'
                content=[]
                for child in elem:
                    if child.tag==qn('w:r'): content.append(run_html(child))
                    elif child.tag==qn('w:hyperlink'):
                        rel=doc.part.rels.get(child.get(qn('r:id')))
                        text=''.join(run_html(r) for r in child if r.tag==qn('w:r'))
                        content.append('<a href="'+html.escape(rel.target_ref,quote=True)+'">'+text+'</a>' if rel else text)
                val='<'+tag+'>'+''.join(content)+'</'+tag+'>'
                num=elem.find('./'+qn('w:pPr')+'/'+qn('w:numPr'))
                if num is not None or st.startswith('List'): val='<ol><li>'+''.join(content)+'</li></ol>' if 'Number' in st else '<ul><li>'+''.join(content)+'</li></ul>'
                output.append(val)
            elif elem.tag==qn('w:tbl'):
                table=[]
                for row in elem.findall(qn('w:tr')):
                    cells=[]
                    for cell in row.findall(qn('w:tc')):
                        span=cell.find('./'+qn('w:tcPr')+'/'+qn('w:gridSpan'))
                        cs=' colspan="'+span.get(qn('w:val'),'1')+'"' if span is not None else ''
                        cells.append('<td'+cs+'>'+blocks(cell)+'</td>')
                    table.append('<tr>'+''.join(cells)+'</tr>')
                output.append('<table>'+''.join(table)+'</table>')
        return ''.join(output)
    return clean(blocks(doc.element.body))
