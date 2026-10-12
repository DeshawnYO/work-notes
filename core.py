"""Local data and project services for the personal workbench."""
from __future__ import annotations

import base64
import ctypes
import hashlib
import io
import json
import os
import re
import shutil
import sqlite3
import sys
import tempfile
import threading
import urllib.error
import urllib.request
import uuid
import zipfile
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from pathlib import Path
from xml.etree import ElementTree as ET

HERE = Path(__file__).resolve().parent
# Ship the PDF reader with the application so uploads also work offline.
if (HERE / 'vendor').is_dir():
    sys.path.insert(0, str(HERE / 'vendor'))
DATA = Path(os.environ.get("PERSONAL_WORKBENCH_DATA_DIR") or (HERE / "data")).resolve()
DATA.mkdir(parents=True,exist_ok=True)
DB_PATH = DATA / "workbench.sqlite3"
PORTABLE_MODE = (HERE / "portable.marker").is_file()
PROJECTS_DIR = HERE / "projects"
ARCHIVE_DIR = "_项目档案"
MAX_READ = 16 * 1024 * 1024
SCAN_LIMIT = 4000
LOCK = threading.RLock()
RUN_LOCK = threading.Lock()

WEEKLY_SYSTEM_PROMPT = """你是一位专业的周报生成助手，为同时担任产品经理和项目经理的用户撰写工作周报。
只能根据提供的本周工作事实写作，不得编造进度、成果、数字、人员或项目。保持提供的项目名称及顺序，不新增或删除项目。
严格按以下结构输出，不添加前言、结语或其他章节：
### 工作周报
【项目名称】
1. 功能开发及优化
  1）事项：完成的工作与当前进度；
2. 项目协同
  1）事项：完成的协调管理工作与当前进度；
对每个项目重复两个章节。功能需求、设计、开发、测试、上线和功能优化放在第一部分；跨团队协调、风险问题处理、资源配置、会议及项目管理放在第二部分。
没有内容的章节保留并写“  1）暂无进展；”。信息不足时写“待补充”，不要猜测；生成的是待用户核对的草稿。
语言专业、简洁、准确。每个事项使用“1）”形式编号，并以中文分号结束。"""


def now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def today() -> str:
    return date.today().isoformat()


def uid() -> str:
    return uuid.uuid4().hex


@contextmanager
def db():
    con = sqlite3.connect(DB_PATH, timeout=30)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    con.execute("PRAGMA busy_timeout=30000")
    try:
        yield con
        con.commit()
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()


def init_db():
    with db() as con:
        con.execute("PRAGMA journal_mode=WAL")
        con.executescript("""
        CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS projects(
          id TEXT PRIMARY KEY,name TEXT NOT NULL,root TEXT NOT NULL UNIQUE,
          background TEXT NOT NULL DEFAULT '',objective TEXT NOT NULL DEFAULT '',
          local_only INTEGER NOT NULL DEFAULT 0,created_at TEXT NOT NULL,
          last_archive_at TEXT,last_archive_digest TEXT,archive_status TEXT DEFAULT '待更新');
        CREATE TABLE IF NOT EXISTS goals(
          id TEXT PRIMARY KEY,title TEXT NOT NULL,project_id TEXT,week_start TEXT NOT NULL,
          status TEXT NOT NULL DEFAULT '进行中',created_at TEXT NOT NULL,
          period_kind TEXT NOT NULL DEFAULT '周',
          FOREIGN KEY(project_id) REFERENCES projects(id));
        CREATE TABLE IF NOT EXISTS tasks(
          id TEXT PRIMARY KEY,title TEXT NOT NULL,details TEXT NOT NULL DEFAULT '',
          project_id TEXT,goal_id TEXT,important INTEGER,urgent INTEGER,
          status TEXT NOT NULL DEFAULT '待安排',planned_date TEXT,due_date TEXT,
          followup_date TEXT,is_interrupt INTEGER NOT NULL DEFAULT 0,
          created_at TEXT NOT NULL,updated_at TEXT NOT NULL,completed_at TEXT,
          FOREIGN KEY(project_id) REFERENCES projects(id),
          FOREIGN KEY(goal_id) REFERENCES goals(id));
        CREATE TABLE IF NOT EXISTS logs(
          id TEXT PRIMARY KEY,body TEXT NOT NULL,project_id TEXT,
          kind TEXT NOT NULL DEFAULT '随手记录',event_date TEXT NOT NULL,
          created_at TEXT NOT NULL,FOREIGN KEY(project_id) REFERENCES projects(id));
        CREATE TABLE IF NOT EXISTS project_items(
          id TEXT PRIMARY KEY,project_id TEXT NOT NULL,kind TEXT NOT NULL,
          title TEXT NOT NULL,details TEXT NOT NULL DEFAULT '',owner TEXT NOT NULL DEFAULT '',
          status TEXT NOT NULL DEFAULT '待确认',due_date TEXT,outcome TEXT NOT NULL DEFAULT '',
          created_at TEXT NOT NULL,updated_at TEXT NOT NULL,closed_at TEXT,
          FOREIGN KEY(project_id) REFERENCES projects(id));
        CREATE TABLE IF NOT EXISTS events(
          id TEXT PRIMARY KEY,subject_type TEXT NOT NULL,subject_id TEXT NOT NULL,
          event_type TEXT NOT NULL,payload TEXT NOT NULL,created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS files(
          project_id TEXT NOT NULL,rel_path TEXT NOT NULL,size INTEGER NOT NULL,
          mtime_ns INTEGER NOT NULL,digest TEXT NOT NULL,excerpt TEXT NOT NULL DEFAULT '',
          extract_status TEXT NOT NULL,first_seen TEXT NOT NULL,last_seen TEXT NOT NULL,
          PRIMARY KEY(project_id,rel_path),
          FOREIGN KEY(project_id) REFERENCES projects(id));
        CREATE TABLE IF NOT EXISTS archive_runs(
          id TEXT PRIMARY KEY,project_id TEXT NOT NULL,input_digest TEXT NOT NULL,
          status TEXT NOT NULL,details TEXT NOT NULL,created_at TEXT NOT NULL,
          FOREIGN KEY(project_id) REFERENCES projects(id));
        CREATE TABLE IF NOT EXISTS reports(
          id TEXT PRIMARY KEY,kind TEXT NOT NULL,period_start TEXT NOT NULL,
          period_end TEXT NOT NULL,body TEXT NOT NULL,created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS model_profiles(
          id TEXT PRIMARY KEY,name TEXT NOT NULL,provider TEXT NOT NULL,
          base_url TEXT NOT NULL,model TEXT NOT NULL,key_blob TEXT NOT NULL DEFAULT '',
          created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS portable_project_paths(
          project_id TEXT PRIMARY KEY,rel_path TEXT NOT NULL,
          FOREIGN KEY(project_id) REFERENCES projects(id));
        CREATE TABLE IF NOT EXISTS ai_archive_files(
          project_id TEXT NOT NULL,rel_path TEXT NOT NULL,size INTEGER NOT NULL,
          mtime_ns INTEGER NOT NULL,extract_status TEXT NOT NULL,
          PRIMARY KEY(project_id,rel_path));
        CREATE TABLE IF NOT EXISTS ai_archive_jobs(
          id TEXT PRIMARY KEY,project_id TEXT NOT NULL,mode TEXT NOT NULL,
          status TEXT NOT NULL,phase TEXT NOT NULL,total_files INTEGER NOT NULL DEFAULT 0,
          selected_files INTEGER NOT NULL DEFAULT 0,processed_files INTEGER NOT NULL DEFAULT 0,
          readable_files INTEGER NOT NULL DEFAULT 0,model_calls INTEGER NOT NULL DEFAULT 0,
          detail TEXT NOT NULL DEFAULT '',created_at TEXT NOT NULL,updated_at TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS idx_tasks_day ON tasks(planned_date,status);
        CREATE TABLE IF NOT EXISTS task_carryovers(
          task_id TEXT NOT NULL,from_date TEXT NOT NULL,to_date TEXT NOT NULL,created_at TEXT NOT NULL,
          PRIMARY KEY(task_id,to_date),FOREIGN KEY(task_id) REFERENCES tasks(id) ON DELETE CASCADE);
        CREATE TABLE IF NOT EXISTS project_dossiers(
          project_id TEXT PRIMARY KEY,template TEXT NOT NULL DEFAULT '',updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS dossier_jobs(
          id TEXT PRIMARY KEY,project_id TEXT NOT NULL,kind TEXT NOT NULL,status TEXT NOT NULL,
          detail TEXT NOT NULL DEFAULT '',created_at TEXT NOT NULL,updated_at TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS idx_logs_project_day ON logs(project_id,event_date);
        CREATE INDEX IF NOT EXISTS idx_files_project ON files(project_id);
        CREATE INDEX IF NOT EXISTS idx_project_items_project ON project_items(project_id,kind,status);
        """)
        # Existing weekly goals remain weekly after the additive migration.
        if "period_kind" not in {row[1] for row in con.execute("PRAGMA table_info(goals)")}:
            con.execute("ALTER TABLE goals ADD COLUMN period_kind TEXT NOT NULL DEFAULT '周'")
        con.execute("INSERT OR IGNORE INTO settings VALUES('archive_hour','21')")
        if "deleted_at" not in {row[1] for row in con.execute("PRAGMA table_info(projects)")}:
            con.execute("ALTER TABLE projects ADD COLUMN deleted_at TEXT")
        if 'source_text' not in {row[1] for row in con.execute('PRAGMA table_info(dossier_jobs)')}:
            con.execute("ALTER TABLE dossier_jobs ADD COLUMN source_text TEXT NOT NULL DEFAULT ''")
        con.execute("UPDATE dossier_jobs SET status='已中断',detail='服务重启，原档案保留，请重新整理',updated_at=? WHERE status='整理中'", (now(),))
        con.execute("UPDATE ai_archive_jobs SET status='已中断',phase='服务重启，需重新加载',updated_at=? WHERE status IN ('排队中','运行中','取消中')", (now(),))
        # Only bundle-owned folders move with the package. Rebind their absolute paths
        # before any scheduled scan or browser request sees the database.
        bundle_root = PROJECTS_DIR.resolve()
        for project_id, rel_path in con.execute("SELECT project_id,rel_path FROM portable_project_paths"):
            rel = Path(rel_path)
            if rel.is_absolute() or ".." in rel.parts or "." in rel.parts:
                continue
            target = (HERE / rel).resolve()
            if bundle_root in target.parents and target.is_dir():
                con.execute("UPDATE projects SET root=? WHERE id=?", (str(target), project_id))
        import v2_data
        v2_data.migrate(con, now())
        personal_root = (DATA / '个人记录' / '未关联项目').resolve()
        for row in con.execute("SELECT id,personal_rel FROM major_records WHERE file_project_id IS NULL AND personal_rel!=''"):
            target = (personal_root / row['personal_rel']).resolve()
            if personal_root in target.parents and target.is_file():
                con.execute('UPDATE major_records SET file_path=? WHERE id=?',(str(target),row['id']))


def rows(sql: str, args=()):
    with db() as con:
        return [dict(r) for r in con.execute(sql, args).fetchall()]


def one(sql: str, args=()):
    with db() as con:
        row = con.execute(sql, args).fetchone()
        return dict(row) if row else None


def setting(key: str, default=""):
    row = one("SELECT value FROM settings WHERE key=?", (key,))
    return row["value"] if row else default


def put_setting(key: str, value: str):
    with db() as con:
        con.execute("INSERT INTO settings VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, value))


def record_event(subject_type: str, subject_id: str, event_type: str, payload: dict):
    with db() as con:
        con.execute("INSERT INTO events VALUES(?,?,?,?,?,?)", (uid(), subject_type, subject_id, event_type, json.dumps(payload, ensure_ascii=False), now()))


def week_start(day: date | None = None) -> str:
    d = day or date.today()
    return (d - timedelta(days=d.weekday())).isoformat()


def validate_day(value, required=False):
    if not value:
        if required:
            raise ValueError("日期不能为空")
        return None
    return date.fromisoformat(value).isoformat()


def create_project(data: dict):
    root = Path(str(data.get("root", "")).strip()).expanduser().resolve()
    if not root.is_dir():
        raise ValueError("请选择已经存在的项目文件夹")
    if root == DATA or DATA in root.parents or root in DATA.parents:
        raise ValueError("项目文件夹不能与工作台数据目录重叠")
    if root.is_symlink():
        raise ValueError("请直接选择真实的项目文件夹")
    existing = one("SELECT * FROM projects WHERE root=?", (str(root),))
    if existing and not existing.get('deleted_at'):
        return existing
    for project in rows("SELECT id,root FROM projects WHERE deleted_at IS NULL"):
        old = Path(project["root"])
        if root in old.parents or old in root.parents:
            raise ValueError("项目文件夹不能相互包含；请分别选择独立目录")
    if existing:
        with db() as con:con.execute('UPDATE projects SET deleted_at=NULL WHERE id=?',(existing['id'],))
        return one('SELECT * FROM projects WHERE id=?',(existing['id'],))
    pid = uid()
    name = str(data.get("name") or root.name).strip()[:120]
    if not name:
        raise ValueError("项目名称不能为空")
    bundle_root = PROJECTS_DIR.resolve()
    if root == bundle_root:
        raise ValueError("请选择 projects 中的单个项目文件夹")
    if PORTABLE_MODE and data.get("copy_into_workspace") and bundle_root not in root.parents:
        if root in bundle_root.parents:
            raise ValueError("不能把整个工作台导入为项目")
        PROJECTS_DIR.mkdir(exist_ok=True)
        safe_name = re.sub(r'[^\w\-\u4e00-\u9fff]+', '-', name).strip('-')[:55] or "项目"
        target = PROJECTS_DIR / f"{pid[:8]}-{safe_name}"
        # Do not pull linked content from outside the selected folder into a copy.
        for base, dirs, files in os.walk(root, followlinks=False):
            if any((Path(base) / child).is_symlink() for child in dirs + files):
                raise ValueError("项目含符号链接，请移除链接后再复制到工具包")
        try:
            shutil.copytree(root, target)
        except Exception:
            if target.parent.resolve() == bundle_root and target.exists():
                shutil.rmtree(target)
            raise
        root = target.resolve()
    with db() as con:
        con.execute("INSERT INTO projects(id,name,root,background,objective,local_only,created_at) VALUES(?,?,?,?,?,?,?)",
                    (pid, name, str(root), str(data.get("background") or ""), str(data.get("objective") or ""), int(bool(data.get("local_only"))), now()))
        if bundle_root in root.parents:
            con.execute("INSERT INTO portable_project_paths VALUES(?,?)", (pid, root.relative_to(HERE).as_posix()))
    archive = root / ARCHIVE_DIR
    if archive.exists() and not archive.is_dir():
        with db() as con:
            con.execute("DELETE FROM portable_project_paths WHERE project_id=?", (pid,))
            con.execute("DELETE FROM projects WHERE id=?", (pid,))
        raise ValueError("项目目录内已有同名文件，无法创建项目档案")
    archive.mkdir(exist_ok=True)
    (archive / "每日进展").mkdir(exist_ok=True)
    (archive / "历史版本").mkdir(exist_ok=True)
    main = archive / "项目档案.md"
    if not main.exists():
        main.write_text(f"# {name}｜项目档案\n\n> 已接入，资料尚待整理。\n", encoding="utf-8")
    scan_project(pid)
    update_archive(pid, force=True)
    return one("SELECT * FROM projects WHERE id=?", (pid,))


def update_project(pid: str, data: dict):
    project = one("SELECT * FROM projects WHERE id=?", (pid,))
    if not project:
        raise FileNotFoundError("项目不存在")
    fields = {k: data[k] for k in ("name", "background", "objective") if k in data}
    if "local_only" in data:
        fields["local_only"] = int(bool(data["local_only"]))
    if fields:
        with db() as con:
            con.execute("UPDATE projects SET " + ",".join(f"{k}=?" for k in fields) + " WHERE id=?", tuple(fields.values()) + (pid,))
    return one("SELECT * FROM projects WHERE id=?", (pid,))


def delete_project(pid: str):
    with RUN_LOCK, db() as con:
        project=con.execute('SELECT * FROM projects WHERE id=? AND deleted_at IS NULL',(pid,)).fetchone()
        if not project:raise FileNotFoundError('项目不存在或已删除')
        if con.execute("SELECT id FROM ai_archive_jobs WHERE project_id=? AND status IN ('排队中','运行中','取消中')",(pid,)).fetchone() or con.execute("SELECT id FROM dossier_jobs WHERE project_id=? AND status='整理中'",(pid,)).fetchone():
            raise ValueError('项目正在整理，请等整理结束后再删除')
        con.execute('UPDATE projects SET deleted_at=? WHERE id=?',(now(),pid))
    return {'deleted':True,'project_id':pid}


def create_goal(data: dict):
    title = str(data.get("title") or "").strip()
    if not title:
        raise ValueError("请输入目标")
    period_kind = data.get("period_kind") or "周"
    if period_kind not in ("周", "月"):
        raise ValueError("目标周期无效")
    anchor = validate_day(data.get("week_start"), True)
    if period_kind == "月" and not anchor.endswith("-01"):
        raise ValueError("月目标日期应为当月第一天")
    pid = data.get("project_id") or None
    if pid and not one("SELECT id FROM projects WHERE id=?", (pid,)):
        raise ValueError("项目不存在")
    gid = uid()
    with db() as con:
        con.execute("INSERT INTO goals(id,title,project_id,week_start,status,created_at,period_kind) VALUES(?,?,?,?,?,?,?)", (gid, title, pid, anchor, "进行中", now(), period_kind))
    return one("SELECT * FROM goals WHERE id=?", (gid,))


def update_goal(gid: str, data: dict):
    if not one("SELECT id FROM goals WHERE id=?", (gid,)):
        raise FileNotFoundError("目标不存在")
    allowed = {k: data[k] for k in ("title", "status", "project_id") if k in data}
    if not allowed:
        return one("SELECT * FROM goals WHERE id=?", (gid,))
    if "title" in allowed:
        allowed["title"] = str(allowed["title"]).strip()
        if not allowed["title"]:
            raise ValueError("请输入目标")
    if "status" in allowed and allowed["status"] not in ("进行中", "已完成", "已取消"):
        raise ValueError("无效的目标状态")
    if "project_id" in allowed:
        allowed["project_id"] = allowed["project_id"] or None
        pid = allowed["project_id"]
        if pid and not one("SELECT id FROM projects WHERE id=?", (pid,)):
            raise ValueError("项目不存在")
        if pid and one("SELECT id FROM tasks WHERE goal_id=? AND project_id IS NOT NULL AND project_id!=? LIMIT 1", (gid, pid)):
            raise ValueError("已有事项属于其他项目，不能直接修改目标所属项目")
    with db() as con:
        con.execute("UPDATE goals SET " + ",".join(f"{k}=?" for k in allowed) + " WHERE id=?", tuple(allowed.values()) + (gid,))
    return one("SELECT * FROM goals WHERE id=?", (gid,))


TASK_FIELDS = ("title", "details", "details_html", "project_id", "goal_id", "important", "urgent", "status", "planned_date", "due_date", "followup_date", "is_interrupt")
STATUSES = ("待安排", "已安排", "进行中", "等待反馈", "受阻", "已完成", "已取消")


def normalize_task(data: dict, current=None):
    values = {k: data[k] for k in TASK_FIELDS if k in data}
    if 'details_html' in values:
        import richtext
        values['details_html'] = richtext.clean(values['details_html'])
        values['details'] = richtext.plain(values['details_html'])
    elif 'details' in values:
        import html
        values['details_html'] = '<p>' + html.escape(str(values['details'] or '')).replace('\n','<br>') + '</p>'
    if "title" in values:
        values["title"] = str(values["title"]).strip()
        if not values["title"]:
            raise ValueError("请输入事项名称")
    if "status" in values and values["status"] not in STATUSES:
        raise ValueError("无效的事项状态")
    for k in ("planned_date", "due_date", "followup_date"):
        if k in values:
            values[k] = validate_day(values[k])
    for k in ("important", "urgent"):
        if k in values:
            values[k] = None if values[k] is None else int(bool(values[k]))
    if "is_interrupt" in values:
        values["is_interrupt"] = int(bool(values["is_interrupt"]))
    for k in ("project_id", "goal_id"):
        if k in values:
            values[k] = values[k] or None
    pid = values.get("project_id", (current or {}).get("project_id"))
    gid = values.get("goal_id", (current or {}).get("goal_id"))
    if gid:
        goal = one("SELECT * FROM goals WHERE id=?", (gid,))
        if not goal:
            raise ValueError("目标不存在")
        if goal["project_id"] and pid and goal["project_id"] != pid:
            raise ValueError("目标和事项属于不同项目")
        if goal["project_id"] and not pid:
            values["project_id"] = goal["project_id"]
    if pid and not one("SELECT id FROM projects WHERE id=?", (pid,)):
        raise ValueError("项目不存在")
    return values


def create_task(data: dict):
    values = normalize_task(data)
    if "title" not in values:
        raise ValueError("请输入事项名称")
    tid = uid()
    stamp = now()
    with db() as con:
        con.execute("INSERT INTO tasks(id,title,details,project_id,goal_id,important,urgent,status,planned_date,due_date,followup_date,is_interrupt,created_at,updated_at,completed_at,details_html) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (tid, values["title"], values.get("details") or "", values.get("project_id"), values.get("goal_id"), values.get("important"), values.get("urgent"), values.get("status") or ("已安排" if values.get("planned_date") else "待安排"), values.get("planned_date"), values.get("due_date"), values.get("followup_date"), values.get("is_interrupt", 0), stamp, stamp, stamp if values.get("status") == "已完成" else None, values.get('details_html','')))
    record_event("task", tid, "创建", values)
    return one("SELECT * FROM tasks WHERE id=?", (tid,))


def update_task(tid: str, data: dict):
    old = one("SELECT * FROM tasks WHERE id=?", (tid,))
    if not old:
        raise FileNotFoundError("事项不存在")
    values = normalize_task(data, old)
    if not values:
        return old
    values["updated_at"] = now()
    if values.get("status") == "已完成" and old["status"] != "已完成":
        values["completed_at"] = now()
    elif values.get("status") and old["status"] == "已完成" and values["status"] != "已完成":
        values["completed_at"] = None
    with db() as con:
        if values.get('status') == '已完成' and con.execute('SELECT id FROM subtasks WHERE task_id=? AND completed=0 LIMIT 1',(tid,)).fetchone():
            raise ValueError('还有未完成的子事项，请先完成子事项再确认父事项')
        con.execute("UPDATE tasks SET " + ",".join(f"{k}=?" for k in values) + " WHERE id=?", tuple(values.values()) + (tid,))
    record_event("task", tid, "更新", {"before": {k: old.get(k) for k in values}, "after": values})
    return one("SELECT * FROM tasks WHERE id=?", (tid,))


def carryover_state():
    target = today()
    previous = (date.fromisoformat(target) - timedelta(days=1)).isoformat()
    confirmed = bool(setting("carryover_confirm_" + target))
    pending = [] if confirmed else rows(
        "SELECT * FROM tasks WHERE status NOT IN ('已完成','已取消') AND "
        "(planned_date=? OR (planned_date IS NULL AND due_date=?)) ORDER BY created_at,id",
        (previous, previous))
    return {"date": target, "previous_date": previous, "confirmed": confirmed, "tasks": pending}


def confirm_carryover(data: dict):
    target = validate_day(data.get("date"))
    if target != today():
        raise ValueError("日期已变化，请回到今天后重新确认昨日事项")
    ids = data.get("task_ids")
    if not isinstance(ids, list) or any(not isinstance(tid, str) for tid in ids):
        raise ValueError("请选择需要加入今天的事项")
    ids = list(dict.fromkeys(ids))
    previous = (date.fromisoformat(target) - timedelta(days=1)).isoformat()
    stamp = now()
    key = "carryover_confirm_" + target
    with db() as con:
        con.execute("BEGIN IMMEDIATE")
        if con.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone():
            return {"confirmed": True, "moved": 0, "already_confirmed": True}
        candidates = {row["id"]: dict(row) for row in con.execute(
            "SELECT * FROM tasks WHERE status NOT IN ('已完成','已取消') AND "
            "(planned_date=? OR (planned_date IS NULL AND due_date=?))", (previous, previous))}
        if any(tid not in candidates for tid in ids):
            raise ValueError("昨日事项已发生变化，请重新查看清单后确认")
        for tid in ids:
            task = candidates[tid]
            con.execute("INSERT INTO task_carryovers VALUES(?,?,?,?)", (tid, previous, target, stamp))
            con.execute("UPDATE tasks SET planned_date=?,updated_at=? WHERE id=?", (target, stamp, tid))
            con.execute("INSERT INTO events VALUES(?,?,?,?,?,?)",
                        (uid(), "task", tid, "顺延", json.dumps({"from_date": previous, "to_date": target,
                         "original_planned_date": task["planned_date"]}, ensure_ascii=False), stamp))
        con.execute("INSERT INTO settings VALUES(?,?)", (key, json.dumps({"task_ids": ids, "confirmed_at": stamp})))
    return {"confirmed": True, "moved": len(ids), "date": target}


def delete_task(tid: str):
    old = one("SELECT * FROM tasks WHERE id=?", (tid,))
    if not old:
        raise FileNotFoundError("事项不存在")
    with db() as con:
        con.execute("DELETE FROM tasks WHERE id=?", (tid,))
    record_event("task", tid, "删除", old)
    return {"deleted": True, "id": tid}


def create_log(data: dict):
    body = str(data.get("body") or "").strip()
    if not body:
        raise ValueError("请输入记录内容")
    pid = data.get("project_id") or None
    if pid and not one("SELECT id FROM projects WHERE id=?", (pid,)):
        raise ValueError("项目不存在")
    lid = uid()
    with db() as con:
        con.execute("INSERT INTO logs VALUES(?,?,?,?,?,?)", (lid, body, pid, str(data.get("kind") or "随手记录"), validate_day(data.get("event_date") or today(), True), now()))
    return one("SELECT * FROM logs WHERE id=?", (lid,))


def update_log(lid: str, data: dict):
    old = one("SELECT * FROM logs WHERE id=?", (lid,))
    if not old:
        raise FileNotFoundError("工作记录不存在")
    values = {k: data[k] for k in ("body", "project_id", "kind", "event_date") if k in data}
    if "body" in values:
        values["body"] = str(values["body"]).strip()
        if not values["body"]:
            raise ValueError("请输入记录内容")
    if "project_id" in values:
        values["project_id"] = values["project_id"] or None
        if values["project_id"] and not one("SELECT id FROM projects WHERE id=?", (values["project_id"],)):
            raise ValueError("项目不存在")
    if "event_date" in values:
        values["event_date"] = validate_day(values["event_date"], True)
    if "kind" in values:
        values["kind"] = str(values["kind"] or "随手记录")
    if values:
        with db() as con:
            con.execute("UPDATE logs SET " + ",".join(f"{k}=?" for k in values) + " WHERE id=?", tuple(values.values()) + (lid,))
        record_event("log", lid, "更新", {"before": {k: old[k] for k in values}, "after": values})
    return one("SELECT * FROM logs WHERE id=?", (lid,))


def month_bounds(month: str) -> tuple[str, str]:
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month or ""):
        raise ValueError("月份格式应为 YYYY-MM")
    start = date.fromisoformat(month + "-01")
    end = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
    return start.isoformat(), end.isoformat()


def month_logs(month: str, project_id: str = "") -> list[dict]:
    start, end = month_bounds(month)
    if project_id and not one("SELECT id FROM projects WHERE id=?", (project_id,)):
        raise ValueError("项目不存在")
    sql = "SELECT * FROM logs WHERE event_date>=? AND event_date<?"
    args = [start, end]
    if project_id:
        sql += " AND project_id=?"
        args.append(project_id)
    return rows(sql + " ORDER BY event_date ASC,created_at ASC,id ASC", args)


def month_task_activity(month: str, project_id: str = "") -> list[dict]:
    start, end = month_bounds(month)
    if project_id and not one("SELECT id FROM projects WHERE id=?", (project_id,)):
        raise ValueError("项目不存在")
    sql = ("SELECT * FROM tasks WHERE status!='已取消' AND "
           "((planned_date>=? AND planned_date<?) OR "
           "(substr(completed_at,1,10)>=? AND substr(completed_at,1,10)<?) OR "
           "(planned_date IS NULL AND substr(created_at,1,10)>=? AND substr(created_at,1,10)<?) OR "
           "id IN (SELECT task_id FROM task_carryovers WHERE (from_date>=? AND from_date<?) OR (to_date>=? AND to_date<?)))")
    args = [start, end] * 5
    if project_id:
        sql += " AND project_id=?"
        args.append(project_id)
    activity = []
    history = {}
    for entry in rows("SELECT task_id,from_date,to_date FROM task_carryovers WHERE (from_date>=? AND from_date<?) OR (to_date>=? AND to_date<?)", (start, end, start, end)):
        history.setdefault(entry["task_id"], []).append(entry)
    for task in rows(sql, args):
        planned = task["planned_date"]
        completed = (task["completed_at"] or "")[:10]
        created = task["created_at"][:10]
        if planned and start <= planned < end:
            activity.append({**task, "event_date": planned,
                             "activity": "完成" if completed == planned else "计划"})
        elif not planned and start <= created < end:
            activity.append({**task, "event_date": created,
                             "activity": "完成" if completed == created else "新增"})
        if completed and start <= completed < end and completed != (planned or created):
            activity.append({**task, "event_date": completed, "activity": "完成"})
        seen = {item["event_date"] for item in activity if item["id"] == task["id"]}
        for entry in history.get(task["id"], []):
            for day in (entry["from_date"], entry["to_date"]):
                if start <= day < end and day not in seen:
                    activity.append({**task, "event_date": day, "activity": "顺延记录"})
                    seen.add(day)
    return sorted(activity, key=lambda item: (item["event_date"], item["created_at"], item["id"]))


def month_logs_markdown(month: str, project_id: str = "") -> str:
    logs = month_logs(month, project_id)
    tasks = month_task_activity(month, project_id)
    projects = {project["id"]: project["name"] for project in rows("SELECT id,name FROM projects")}
    days = sorted({entry["event_date"] for entry in logs + tasks})
    heading = f"# {month[:4]} 年 {month[5:]} 月工作记录"
    lines = [heading, "", f"有内容的日期：{len(days)} 天 · 工作记录：{len(logs)} 条 · 事项动态：{len(tasks)} 项",
             f"项目范围：{projects.get(project_id, '全部项目') if project_id else '全部项目'}", ""]
    for day in days:
        weekday = "一二三四五六日"[date.fromisoformat(day).weekday()]
        lines += [f"## {day} 星期{weekday}", ""]
        for entry in (item for item in logs if item["event_date"] == day):
            name = projects.get(entry["project_id"], "未关联项目")
            lines.append(f"- **{entry['kind']}** · {name}")
            lines.extend(f"  > {line}" for line in entry["body"].splitlines())
            lines.append("")
        daily_tasks = [item for item in tasks if item["event_date"] == day]
        if daily_tasks:
            lines += ["### 当天事项", ""]
            for task in daily_tasks:
                name = projects.get(task["project_id"], "未关联项目")
                lines.append(f"- **{task['activity']}** · {task['title']}（{name}；当前状态：{task['status']}）")
                if task["details"].strip():
                    lines.extend(f"  > {line}" for line in task["details"].splitlines())
                lines.append("")
    if not days:
        lines += ["本月暂无工作记录。", ""]
    return "\n".join(lines)


def browse_folders(path: str = "") -> dict:
    folder = Path(path).expanduser() if path else Path.home()
    if not folder.is_dir():
        raise ValueError("文件夹不存在，请选择其他位置")
    folder = folder.resolve()
    try:
        children = sorted((item for item in folder.iterdir() if item.is_dir() and not item.is_symlink()),
                          key=lambda item: (item.name.startswith("."), item.name.casefold()))
    except PermissionError as exc:
        raise ValueError("没有权限读取这个文件夹，请选择其他位置") from exc
    drives = [f"{letter}:\\" for letter in "ABCDEFGHIJKLMNOPQRSTUVWXYZ" if Path(f"{letter}:\\").is_dir()] if os.name == "nt" else ["/"]
    return {"path": str(folder), "parent": str(folder.parent) if folder.parent != folder else "",
            "folders": [{"name": item.name, "path": str(item)} for item in children], "drives": drives}


PROJECT_ITEM_KINDS = ("里程碑", "风险", "问题", "决策", "跟进")
PROJECT_ITEM_STATUSES = ("待确认", "进行中", "等待反馈", "受阻", "已完成", "已关闭", "已取消")


def normalize_project_item(data: dict, current=None):
    values = {k: data[k] for k in ("project_id", "kind", "title", "details", "owner", "status", "due_date", "outcome") if k in data}
    if "project_id" in values:
        if not one("SELECT id FROM projects WHERE id=?", (values["project_id"],)):
            raise ValueError("项目不存在")
    if "kind" in values and values["kind"] not in PROJECT_ITEM_KINDS:
        raise ValueError("台账类型不支持")
    if "status" in values and values["status"] not in PROJECT_ITEM_STATUSES:
        raise ValueError("台账状态不支持")
    if "title" in values:
        values["title"] = str(values["title"]).strip()
        if not values["title"]:
            raise ValueError("请输入标题")
    for key in ("details", "owner", "outcome"):
        if key in values:
            values[key] = str(values[key] or "").strip()
    if "due_date" in values:
        values["due_date"] = validate_day(values["due_date"])
    if current and "project_id" in values and values["project_id"] != current["project_id"]:
        raise ValueError("不能将台账条目直接移到其他项目")
    return values


def create_project_item(data: dict):
    values = normalize_project_item(data)
    if not all(values.get(k) for k in ("project_id", "kind", "title")):
        raise ValueError("请选择项目、类型并填写标题")
    stamp = now()
    item_id = uid()
    status = values.get("status") or "待确认"
    closed_at = stamp if status in ("已完成", "已关闭", "已取消") else None
    with db() as con:
        con.execute("INSERT INTO project_items VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                    (item_id, values["project_id"], values["kind"], values["title"], values.get("details") or "",
                     values.get("owner") or "", status, values.get("due_date"), values.get("outcome") or "", stamp, stamp, closed_at))
    record_event("project_item", item_id, "创建", values)
    return one("SELECT * FROM project_items WHERE id=?", (item_id,))


def update_project_item(item_id: str, data: dict):
    old = one("SELECT * FROM project_items WHERE id=?", (item_id,))
    if not old:
        raise FileNotFoundError("项目台账条目不存在")
    values = normalize_project_item(data, old)
    if not values:
        return old
    stamp = now()
    values["updated_at"] = stamp
    if "status" in values:
        values["closed_at"] = stamp if values["status"] in ("已完成", "已关闭", "已取消") and old["status"] not in ("已完成", "已关闭", "已取消") else (None if values["status"] not in ("已完成", "已关闭", "已取消") else old["closed_at"])
    with db() as con:
        con.execute("UPDATE project_items SET " + ",".join(f"{k}=?" for k in values) + " WHERE id=?", tuple(values.values()) + (item_id,))
    record_event("project_item", item_id, "更新", {"before": {k: old.get(k) for k in values}, "after": values})
    return one("SELECT * FROM project_items WHERE id=?", (item_id,))


def safe_path(pid: str, rel: str, allow_new=False) -> tuple[dict, Path]:
    project = one("SELECT * FROM projects WHERE id=?", (pid,))
    if not project:
        raise FileNotFoundError("项目不存在")
    root = Path(project["root"]).resolve()
    if not root.is_dir():
        raise FileNotFoundError("项目文件夹暂不可用，请重新连接")
    rel = str(rel or "").replace("\\", "/")
    parts = Path(rel).parts
    if not parts or Path(rel).is_absolute() or rel.startswith("/") or any(p in ("..", ".", "") for p in parts):
        raise ValueError("无效的项目内路径")
    if parts[0] == ARCHIVE_DIR:
        raise ValueError("项目档案请在档案页面管理")
    target = root.joinpath(*parts)
    real = target.resolve(strict=False)
    if root not in real.parents:
        raise ValueError("不能访问项目文件夹之外的路径")
    if target.is_symlink():
        raise ValueError("不自动操作快捷方式或链接目标")
    if not allow_new and not target.exists():
        raise FileNotFoundError("文件不存在")
    return project, target


def file_tree(pid: str, subdir=""):
    project = one("SELECT * FROM projects WHERE id=?", (pid,))
    if not project:
        raise FileNotFoundError("项目不存在")
    root = Path(project["root"])
    base = root if not subdir else safe_path(pid, subdir)[1]
    if not base.is_dir():
        raise ValueError("请选择文件夹")
    result = []
    for item in sorted(base.iterdir(), key=lambda x: (not x.is_dir(), x.name.lower())):
        if item.name == ARCHIVE_DIR or item.name.startswith("~$") or item.name.startswith("."):
            continue
        rel = item.relative_to(root).as_posix()
        linked = item.is_symlink() or root.resolve() not in item.resolve(strict=False).parents
        stat = item.stat() if item.exists() else None
        result.append({"name": item.name, "path": rel, "directory": item.is_dir(), "linked": linked,
                       "size": stat.st_size if stat else 0, "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="minutes") if stat else ""})
    return result


def extract_text(path: Path, char_limit: int | None = 40000) -> tuple[str, str]:
    if path.stat().st_size > MAX_READ:
        return "", "超过 16 MB，未提取正文"
    ext = path.suffix.lower()
    def limited(body: str) -> str:
        return body[:char_limit] if char_limit is not None else body
    try:
        if ext in (".md", ".txt", ".csv", ".json", ".html"):
            raw = path.read_bytes()
            for enc in ("utf-8-sig", "gb18030", "utf-16"):
                try:
                    return limited(raw.decode(enc)), "已读取文本"
                except UnicodeDecodeError:
                    pass
            return "", "无法识别文本编码"
        if ext == ".docx":
            with zipfile.ZipFile(path) as z:
                root = ET.fromstring(z.read("word/document.xml"))
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                paras = []
                for p in root.findall(".//w:p", ns):
                    content = "".join((n.text or "") for n in p.findall(".//w:t", ns)).strip()
                    if content:
                        paras.append(content)
                return limited("\n".join(paras)), "已读取正文和可提取表格文字；图片、批注等未覆盖"
        if ext == ".pptx":
            with zipfile.ZipFile(path) as z:
                slides = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                                key=lambda n: int(re.search(r"slide(\d+)", n).group(1)))
                chunks = []
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                for i, name in enumerate(slides, 1):
                    root = ET.fromstring(z.read(name))
                    text = " ".join((n.text or "") for n in root.findall(".//a:t", ns)).strip()
                    if text:
                        chunks.append(f"第{i}页：{text}")
                return limited("\n".join(chunks)), "已读取幻灯片文字；图表、备注和图片未覆盖"
        if ext == ".xlsx":
            try:
                from openpyxl import load_workbook
            except ImportError:
                return "", "缺少 openpyxl，未提取表格正文"
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                chunks = []
                for sheet in workbook.worksheets:
                    chunks.append(f"工作表：{sheet.title}")
                    for row in sheet.iter_rows(values_only=True):
                        values = [str(value).replace("\n", " ") for value in row if value is not None]
                        if values:
                            chunks.append(" | ".join(values))
                return limited("\n".join(chunks)), "已读取单元格值；公式和图表未覆盖"
            finally:
                workbook.close()
        if ext == ".xls":
            try:
                import xlrd
            except ImportError:
                return "", "缺少 xlrd，未提取旧版表格正文"
            book = xlrd.open_workbook(path, on_demand=True)
            try:
                chunks = []
                for sheet in book.sheets():
                    chunks.append(f"工作表：{sheet.name}")
                    for index in range(sheet.nrows):
                        values = [str(value).replace("\n", " ") for value in sheet.row_values(index) if value != ""]
                        if values:
                            chunks.append(" | ".join(values))
                return limited("\n".join(chunks)), "已读取单元格值；公式和图表未覆盖"
            finally:
                book.release_resources()
        if ext == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                return "", "缺少 pypdf，未提取 PDF 正文"
            reader = PdfReader(str(path))
            return limited("\n".join(page.extract_text() or "" for page in reader.pages)), "已读取 PDF 可提取文字；扫描图片未覆盖"
        return "", "此格式尚不支持内容提取，可用本机应用打开"
    except (OSError, KeyError, ValueError, zipfile.BadZipFile, ET.ParseError) as exc:
        return "", f"读取失败：{type(exc).__name__}"


def read_file(pid: str, rel: str):
    _, path = safe_path(pid, rel)
    if not path.is_file():
        raise ValueError("请选择文件")
    body, status = extract_text(path)
    editable = path.suffix.lower() in (".md", ".txt") and status == "已读取文本"
    return {"path": rel, "body": body, "status": status, "editable": editable, "modified_ns": str(path.stat().st_mtime_ns)}


def atomic_write(path: Path, content: bytes):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=".workbench-", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as out:
            out.write(content)
            out.flush()
            os.fsync(out.fileno())
        os.replace(name, path)
    finally:
        if os.path.exists(name):
            os.unlink(name)


def save_file(pid: str, rel: str, content: bytes, expected_ns=None, overwrite=False):
    project, path = safe_path(pid, rel, allow_new=True)
    if path.exists() and path.is_dir():
        raise ValueError("同名文件夹已存在")
    if path.exists() and expected_ns is not None and path.stat().st_mtime_ns != int(expected_ns):
        raise ValueError("文件已在其他窗口修改，请重新载入或另存副本")
    if path.exists() and expected_ns is None and not overwrite:
        raise ValueError("同名文件已存在，请更换名称或明确覆盖")
    if path.parent.resolve() != Path(project["root"]).resolve() and Path(project["root"]).resolve() not in path.parent.resolve().parents:
        raise ValueError("不能写入项目目录之外")
    atomic_write(path, content)
    record_event("file", pid, "保存", {"path": rel, "size": len(content)})
    scan_project(pid)
    return {"path": rel, "modified_ns": str(path.stat().st_mtime_ns)}


def create_folder(pid: str, rel: str):
    _, path = safe_path(pid, rel, allow_new=True)
    if path.exists():
        raise ValueError("文件或文件夹已存在")
    path.mkdir(parents=True)
    return {"path": rel}


def create_document(pid: str, rel: str):
    from io import BytesIO
    ext = Path(rel).suffix.lower()
    if ext in (".md", ".txt"):
        content = (f"# {Path(rel).stem}\n" if ext == ".md" else "").encode("utf-8")
    elif ext == ".docx":
        from docx import Document
        document = Document()
        document.add_heading(Path(rel).stem, 0)
        out = BytesIO()
        document.save(out)
        content = out.getvalue()
    elif ext == ".pptx":
        from pptx import Presentation
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = Path(rel).stem
        out = BytesIO()
        presentation.save(out)
        content = out.getvalue()
    else:
        raise ValueError("可新建 Markdown、TXT、Word 或 PPT 文件")
    return save_file(pid, rel, content)


def rename_file(pid: str, old_rel: str, new_rel: str):
    _, source = safe_path(pid, old_rel)
    _, target = safe_path(pid, new_rel, allow_new=True)
    if target.exists():
        raise ValueError("目标名称已存在")
    if source == target or source in target.parents:
        raise ValueError("不能把文件夹移到自身内部")
    source.rename(target)
    record_event("file", pid, "重命名", {"old": old_rel, "new": new_rel})
    scan_project(pid)
    return {"path": new_rel}


def delete_file(pid: str, rel: str):
    _, path = safe_path(pid, rel)
    if os.name != "nt":
        raise RuntimeError("可恢复删除仅支持 Windows")
    class FileOp(ctypes.Structure):
        _fields_ = [("hwnd", ctypes.c_void_p), ("wFunc", ctypes.c_uint),
                    ("pFrom", ctypes.c_wchar_p), ("pTo", ctypes.c_wchar_p),
                    ("fFlags", ctypes.c_ushort), ("fAnyOperationsAborted", ctypes.c_int),
                    ("hNameMappings", ctypes.c_void_p), ("lpszProgressTitle", ctypes.c_wchar_p)]
    path_buffer = ctypes.create_unicode_buffer(str(path) + "\0\0")
    operation = FileOp(None, 3, ctypes.cast(path_buffer, ctypes.c_wchar_p), None,
                       0x0040 | 0x0010 | 0x0004 | 0x0400, 0, None, None)
    result = ctypes.windll.shell32.SHFileOperationW(ctypes.byref(operation))
    if result or operation.fAnyOperationsAborted:
        raise RuntimeError("未能移入系统回收站，原文件保持不变")
    record_event("file", pid, "移入回收站", {"path": rel})
    scan_project(pid)
    return {"deleted": rel, "recoverable": True}


def open_native(pid: str, rel: str):
    _, path = safe_path(pid, rel)
    if os.name != "nt":
        raise RuntimeError("当前只支持 Windows 原生打开")
    os.startfile(str(path))
    return {"opened": rel}


def scan_project(pid: str):
    project = one("SELECT * FROM projects WHERE id=?", (pid,))
    if not project:
        raise FileNotFoundError("项目不存在")
    root = Path(project["root"])
    if not root.is_dir():
        raise FileNotFoundError("项目文件夹暂不可用")
    initial_index = not one("SELECT id FROM archive_runs WHERE project_id=? LIMIT 1", (pid,)) and not one("SELECT rel_path FROM files WHERE project_id=? LIMIT 1", (pid,))
    found = set()
    changes = []
    count = 0
    for base, dirs, files in os.walk(root, followlinks=False):
        dirs[:] = [d for d in dirs if d != ARCHIVE_DIR and not d.startswith(".") and not (Path(base) / d).is_symlink()]
        for name in files:
            if name.startswith(("~$", ".workbench-", ".")):
                continue
            path = Path(base) / name
            if path.is_symlink():
                continue
            rel = path.relative_to(root).as_posix()
            found.add(rel)
            count += 1
            if count > SCAN_LIMIT:
                return {"changes": changes, "status": "文件超过扫描上限，需缩小项目目录", "count": count}
            try:
                stat = path.stat()
                old = one("SELECT * FROM files WHERE project_id=? AND rel_path=?", (pid, rel))
                if old and old["size"] == stat.st_size and old["mtime_ns"] == stat.st_mtime_ns:
                    continue
                if stat.st_size > MAX_READ:
                    digest = f"large:{stat.st_size}:{stat.st_mtime_ns}"
                    excerpt, status = "", "文件过大，未提取"
                else:
                    digest = hashlib.sha256(path.read_bytes()).hexdigest()
                    if old and old["digest"] == digest:
                        with db() as con:
                            con.execute("UPDATE files SET mtime_ns=?,last_seen=? WHERE project_id=? AND rel_path=?", (stat.st_mtime_ns, now(), pid, rel))
                        continue
                    excerpt, status = extract_text(path)
                stamp = now()
                with db() as con:
                    con.execute("INSERT INTO files VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(project_id,rel_path) DO UPDATE SET size=excluded.size,mtime_ns=excluded.mtime_ns,digest=excluded.digest,excerpt=excluded.excerpt,extract_status=excluded.extract_status,last_seen=excluded.last_seen",
                                (pid, rel, stat.st_size, stat.st_mtime_ns, digest, excerpt[:10000], status, stamp, stamp))
                changes.append({"path": rel, "kind": "新增" if not old else "修改", "status": status})
            except (OSError, PermissionError):
                changes.append({"path": rel, "kind": "读取失败", "status": "文件暂不可读"})
    indexed = rows("SELECT rel_path FROM files WHERE project_id=?", (pid,))
    for item in indexed:
        if item["rel_path"] not in found:
            with db() as con:
                con.execute("DELETE FROM files WHERE project_id=? AND rel_path=?", (pid, item["rel_path"]))
            changes.append({"path": item["rel_path"], "kind": "移出或删除", "status": "原文件不在项目目录中"})
    if changes and not initial_index:
        record_event("project", pid, "文件变化", {"changes": changes[:100], "total": len(changes)})
    return {"changes": changes, "status": "完成", "count": count}


def digest_text(data: str):
    return hashlib.sha256(data.encode("utf-8")).hexdigest()


def archive_material(pid: str):
    project = one("SELECT * FROM projects WHERE id=?", (pid,))
    if not project:
        raise FileNotFoundError("项目不存在")
    all_files = rows("SELECT rel_path,digest,extract_status,last_seen FROM files WHERE project_id=? ORDER BY rel_path", (pid,))
    files = rows("SELECT rel_path,digest,excerpt,extract_status,last_seen FROM files WHERE project_id=? ORDER BY last_seen DESC,rel_path LIMIT 30", (pid,))
    logs = rows("SELECT id,body,kind,event_date FROM logs WHERE project_id=? ORDER BY event_date DESC,created_at DESC LIMIT 30", (pid,))
    tasks = rows("SELECT id,title,status,planned_date,completed_at FROM tasks WHERE project_id=? AND status!='已取消' ORDER BY updated_at DESC LIMIT 30", (pid,))
    items = rows("SELECT id,kind,title,details,owner,status,due_date,outcome,updated_at FROM project_items WHERE project_id=? AND status!='已取消' ORDER BY updated_at DESC LIMIT 100", (pid,))
    payload = {"name": project["name"], "background": project["background"], "objective": project["objective"],
               "files": files, "file_index": all_files, "logs": logs, "tasks": tasks, "items": items}
    return project, payload, digest_text(json.dumps(payload, ensure_ascii=False, sort_keys=True))


def model_summary(prompt: str, project=None):
    profile_id = setting("active_model", "")
    if not profile_id:
        return "", "未配置模型，已用本地事实整理"
    profile = one("SELECT * FROM model_profiles WHERE id=?", (profile_id,))
    if not profile:
        return "", "模型配置不存在，已用本地事实整理"
    if project and project["local_only"] and profile["provider"] != "ollama":
        return "", "项目限制为仅本地处理，已用本地事实整理"
    try:
        return call_model(profile, prompt), "已由模型生成待核实整理"
    except (OSError, ValueError, RuntimeError, urllib.error.URLError, TimeoutError) as exc:
        return "", f"模型不可用（{type(exc).__name__}），已用本地事实整理"


def archive_markdown(project: dict, data: dict, summary: str, note: str):
    lines = [f"# {project['name']}｜项目档案", "", f"> 更新于 {now()}。{note}。", "",
             "## 背景与目标", "", f"背景：{data['background'] or '待补充。'}", "", f"目标：{data['objective'] or '待补充。'}", "",
             "## 本期工作记录", ""]
    if data["logs"]:
        lines += [f"- {l['event_date']}｜{l['kind']}｜{l['body']} （工作记录 {l['id'][:8]}）" for l in data["logs"]]
    else:
        lines.append("暂无已关联项目的工作记录。")
    lines += ["", "## 事项与后续行动", ""]
    if data["tasks"]:
        lines += [f"- [{t['status']}] {t['title']}" + (f"，计划 {t['planned_date']}" if t["planned_date"] else "") for t in data["tasks"]]
    else:
        lines.append("暂无已关联事项。")
    lines += ["", "## 项目推进台账", ""]
    if data["items"]:
        for kind in PROJECT_ITEM_KINDS:
            subset = [item for item in data["items"] if item["kind"] == kind]
            if not subset:
                continue
            lines += [f"### {kind}", ""]
            for item in subset:
                note = f"；负责人/跟进人：{item['owner']}" if item["owner"] else ""
                due = f"；目标日期：{item['due_date']}" if item["due_date"] else ""
                outcome = f"；结果：{item['outcome']}" if item["outcome"] else ""
                lines.append(f"- [{item['status']}] {item['title']}{note}{due}{outcome}（台账 {item['id'][:8]}）")
            lines.append("")
    else:
        lines.append("暂无项目推进台账条目。")
    lines += ["", "## 项目文件与提取范围", "", f"已索引 {len(data['file_index'])} 个文件；以下展示最近变化的 {len(data['files'])} 个文件片段。Word、PPT 的图片、图表和批注等未提取；AI 建议仅基于输入长度允许的部分材料。", ""]
    if data["files"]:
        for f in data["files"]:
            excerpt = re.sub(r"\s+", " ", f["excerpt"][:280]).strip()
            lines.append(f"- `{f['rel_path']}`｜{f['extract_status']}｜版本 {f['digest'][:10]}" + (f"｜片段：{excerpt}" if excerpt else ""))
    else:
        lines.append("暂无已索引项目文件。")
    if len(data["file_index"]) > len(data["files"]):
        lines += ["", "### 其余已索引文件", ""]
        recent = {f["rel_path"] for f in data["files"]}
        lines += [f"- `{f['rel_path']}`｜{f['extract_status']}｜版本 {f['digest'][:10]}" for f in data["file_index"] if f["rel_path"] not in recent]
    if summary:
        lines += ["", "## AI 整理建议（待本人核实）", "", summary.strip()]
    lines += ["", "## 人工补充", "", "请在同目录的《人工补充.md》中记录人工确认和长期说明。", ""]
    return "\n".join(lines)


def update_archive(pid: str, force=False):
    with RUN_LOCK:
        if one('SELECT project_id FROM project_dossiers WHERE project_id=?',(pid,)):
            project=one('SELECT * FROM projects WHERE id=? AND deleted_at IS NULL',(pid,))
            if not project:raise FileNotFoundError('项目不存在')
            daily=Path(project['root'])/ARCHIVE_DIR/'每日进展'/f'{today()}.md'
            logs=rows('SELECT kind,body FROM logs WHERE project_id=? AND event_date=? ORDER BY created_at',(pid,today()))
            atomic_write(daily,(f"# {today()}｜{project['name']}\n\n"+'\n'.join(f"- {x['kind']}：{x['body']}" for x in logs)+'\n').encode('utf-8'))
            return {'status':'手动档案保留，已更新每日进展','project_id':pid,'last_archive_at':project['last_archive_at']}
        ai_run = one("SELECT id FROM ai_archive_jobs WHERE project_id=? AND status='成功' AND mode='full' LIMIT 1", (pid,))
        if ai_run:
            project = one("SELECT * FROM projects WHERE id=?", (pid,))
            if not project:
                raise FileNotFoundError("项目不存在")
            if (Path(project["root"]) / ARCHIVE_DIR / "项目档案.md").is_file():
                daily = Path(project["root"]) / ARCHIVE_DIR / "每日进展" / f"{today()}.md"
                today_logs = rows("SELECT kind,body FROM logs WHERE project_id=? AND event_date=? ORDER BY created_at", (pid, today()))
                atomic_write(daily, (f"# {today()}｜{project['name']}\n\n" +
                                     "\n".join(f"- {item['kind']}：{item['body']}" for item in today_logs) + "\n").encode("utf-8"))
                return {"status": "AI档案保留，已更新每日进展", "project_id": pid, "last_archive_at": project["last_archive_at"]}
        scan = scan_project(pid)
        if scan["status"] != "完成":
            raise RuntimeError(scan["status"])
        project, data, input_digest = archive_material(pid)
        if not force:
            last = one("SELECT * FROM archive_runs WHERE project_id=? AND input_digest=? AND status='成功' ORDER BY created_at DESC LIMIT 1", (pid, input_digest))
            if last:
                return {"status": "无新增", "project_id": pid, "last_archive_at": project["last_archive_at"]}
        archive = Path(project["root"]) / ARCHIVE_DIR
        archive.mkdir(exist_ok=True)
        main = archive / "项目档案.md"
        if main.exists() and project["last_archive_digest"]:
            current = hashlib.sha256(main.read_bytes()).hexdigest()
            if current != project["last_archive_digest"]:
                target = archive / f"待合并-{datetime.now():%Y%m%d-%H%M%S}.md"
                summary, note = model_summary("请仅根据资料提取项目最新进展，不要补造事实。\n" + json.dumps(data, ensure_ascii=False)[:14000], project)
                atomic_write(target, archive_markdown(project, data, summary, note).encode("utf-8"))
                with db() as con:
                    con.execute("UPDATE projects SET archive_status='待合并' WHERE id=?", (pid,))
                    con.execute("INSERT INTO archive_runs VALUES(?,?,?,?,?,?)", (uid(), pid, input_digest, "待合并", str(target), now()))
                return {"status": "待合并", "candidate": str(target)}
        prompt = "你是项目记录助手。只根据以下材料，简明整理最新进展、实际成果、风险与待办；不能声称未证实的验收或收益；每条结论指出文件路径或记录标识。\n" + json.dumps(data, ensure_ascii=False)[:14000]
        summary, note = model_summary(prompt, project)
        body = archive_markdown(project, data, summary, note)
        if main.exists():
            historical = archive / "历史版本" / f"{datetime.now():%Y%m%d-%H%M%S}-{uid()[:6]}.md"
            shutil.copy2(main, historical)
        atomic_write(main, body.encode("utf-8"))
        archive_digest = hashlib.sha256(main.read_bytes()).hexdigest()
        with db() as con:
            con.execute("UPDATE projects SET last_archive_at=?,last_archive_digest=?,archive_status='成功' WHERE id=?", (now(), archive_digest, pid))
            con.execute("INSERT INTO archive_runs VALUES(?,?,?,?,?,?)", (uid(), pid, input_digest, "成功", note, now()))
        daily = archive / "每日进展" / f"{today()}.md"
        atomic_write(daily, (f"# {today()}｜{project['name']}\n\n" + "\n".join(f"- {l['kind']}：{l['body']}" for l in data["logs"] if l["event_date"] == today()) + "\n").encode("utf-8"))
        return {"status": "成功", "project_id": pid, "last_archive_at": now(), "files_changed": len(scan["changes"]), "note": note}


def report_period(kind: str, anchor: str | None = None):
    d = date.fromisoformat(anchor) if anchor else date.today()
    if kind == "周报":
        start = d - timedelta(days=d.weekday())
        end = start + timedelta(days=7)
    elif kind == "月报":
        start = d.replace(day=1)
        end = (start.replace(year=start.year + 1, month=1) if start.month == 12 else start.replace(month=start.month + 1))
    else:
        raise ValueError("报告类型不支持")
    return start.isoformat(), end.isoformat()


def weekly_report_body(project_names: list[str], logs: list[dict], tasks: list[dict], items: list[dict]) -> str:
    """Create a usable weekly draft even when no model is configured."""
    development_terms = ("功能", "需求", "设计", "原型", "开发", "优化", "研发", "上线", "测试", "报表", "页面", "接口", "修复", "迭代", "方案", "评审")
    collaboration_terms = ("协调", "沟通", "跟进", "对接", "会议", "风险", "问题", "排查", "配置", "迁移", "合同", "商务", "备案", "汇报", "推进", "资源", "人员", "供应商")

    def category(text: str, kind: str = "") -> int:
        if kind in ("风险", "问题", "决策", "跟进", "会议纪要", "决策与变更"):
            return 1
        if any(term in text for term in development_terms):
            return 0
        if any(term in text for term in collaboration_terms):
            return 1
        return 0 if kind in ("工作进展", "结果记录", "里程碑") else 1

    grouped = {name: [[], []] for name in project_names}
    for entry in logs:
        name = entry["project_name"] or "未归属"
        if name in grouped:
            body = str(entry["body"]).strip().rstrip("；;。")
            grouped[name][category(body, entry["kind"])].append(body + "；")
    for entry in tasks:
        name = entry["project_name"] or "未归属"
        if name not in grouped or entry["status"] == "已取消":
            continue
        title = str(entry["title"]).strip().rstrip("；;。")
        details = str(entry.get("details") or "").strip().rstrip("；;。")
        status = "已完成" if entry["status"] == "已完成" else f"当前{entry['status']}"
        description = f"{title}，{status}" + (f"；{details}" if details else "") + "；"
        grouped[name][category(title + details)].append(description)
    for entry in items:
        name = entry["project_name"]
        if name not in grouped or entry["status"] == "已取消":
            continue
        title = str(entry["title"]).strip().rstrip("；;。")
        outcome = str(entry.get("outcome") or "").strip().rstrip("；;。")
        description = f"{title}，当前{entry['status']}" + (f"；{outcome}" if outcome else "") + "；"
        grouped[name][category(title + outcome, entry["kind"])].append(description)

    lines = ["### 工作周报", ""]
    for name in project_names:
        lines += [f"【{name}】", "", "1. 功能开发及优化"]
        lines += [f"  {index}）{entry}" for index, entry in enumerate(grouped[name][0], 1)] or ["  1）暂无进展；"]
        lines += ["", "2. 项目协同"]
        lines += [f"  {index}）{entry}" for index, entry in enumerate(grouped[name][1], 1)] or ["  1）暂无进展；"]
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def valid_weekly_report(body: str, project_names: list[str]) -> bool:
    if not body.startswith("### 工作周报\n"):
        return False
    if re.findall(r"^【(.+)】$", body, re.M) != project_names:
        return False
    parts = re.split(r"^【.+】$", body, flags=re.M)[1:]
    if len(parts) != len(project_names):
        return False
    for part in parts:
        if part.count("1. 功能开发及优化") != 1 or part.count("2. 项目协同") != 1:
            return False
        before, after = part.split("1. 功能开发及优化", 1)[1].split("2. 项目协同", 1)
        if not any(re.fullmatch(r"\d+）.+；", line.strip()) for line in before.splitlines()):
            return False
        if not any(re.fullmatch(r"\d+）.+；", line.strip()) for line in after.splitlines()):
            return False
        for section in part.splitlines():
            stripped = section.strip()
            if not stripped or stripped in ("1. 功能开发及优化", "2. 项目协同"):
                continue
            if not re.fullmatch(r"\d+）.+；", stripped):
                return False
    return True


def generate_report(kind: str, anchor=None, automatic=False):
    start, end = report_period(kind, anchor)
    if automatic:
        existing = one("SELECT * FROM reports WHERE kind=? AND period_start=? AND period_end=? ORDER BY created_at DESC LIMIT 1", (kind, start, end))
        if existing:
            return existing
    logs = rows("SELECT l.*,p.name AS project_name,p.local_only FROM logs l LEFT JOIN projects p ON p.id=l.project_id WHERE l.event_date>=? AND l.event_date<? ORDER BY l.event_date,l.created_at", (start, end))
    goals = rows("SELECT g.*,p.name AS project_name,p.local_only FROM goals g LEFT JOIN projects p ON p.id=g.project_id WHERE g.period_kind=? AND g.week_start=? AND g.status!='已取消' ORDER BY g.created_at", ("周" if kind == "周报" else "月", start))
    tasks = rows("SELECT t.*,p.name AS project_name,p.local_only FROM tasks t LEFT JOIN projects p ON p.id=t.project_id WHERE (t.completed_at>=? AND t.completed_at<?) OR (t.status NOT IN ('已完成','已取消') AND substr(t.created_at,1,10)<?) ORDER BY t.updated_at", (start, end, end))
    items = rows("SELECT i.*,p.name AS project_name,p.local_only FROM project_items i JOIN projects p ON p.id=i.project_id WHERE (substr(i.updated_at,1,10)>=? AND substr(i.updated_at,1,10)<?) OR (i.status NOT IN ('已完成','已关闭','已取消') AND substr(i.created_at,1,10)<?) ORDER BY i.updated_at", (start, end, end))
    file_events = rows("SELECT e.*,p.name AS project_name FROM events e JOIN projects p ON p.id=e.subject_id WHERE e.subject_type='project' AND e.event_type='文件变化' AND substr(e.created_at,1,10)>=? AND substr(e.created_at,1,10)<? ORDER BY e.created_at", (start, end))
    if automatic and not logs and not goals and not tasks and not items and not file_events:
        return None
    if kind == "周报":
        weekly_tasks = [entry for entry in tasks if entry["status"] != "已取消" and
                        (start <= entry["updated_at"][:10] < end or
                         start <= (entry["planned_date"] or "") < end or
                         start <= (entry["completed_at"] or "")[:10] < end)]
        weekly_items = [entry for entry in items if entry["status"] != "已取消" and start <= entry["updated_at"][:10] < end]
        report_projects = rows("SELECT name,local_only FROM projects WHERE deleted_at IS NULL ORDER BY created_at,id")
        project_names = [project["name"] for project in report_projects]
        if any(not entry["project_name"] for entry in logs + weekly_tasks):
            project_names.append("未归属")
        if not project_names:
            project_names = ["未归属"]
        body = weekly_report_body(project_names, logs, weekly_tasks, weekly_items)
        profile = one("SELECT * FROM model_profiles WHERE id=?", (setting("active_model"),))
        source = logs + weekly_tasks + weekly_items + goals
        if profile and (profile["provider"] == "ollama" or not any(entry["local_only"] for entry in source + report_projects)):
            facts = {"projects": project_names, "period": [start, end], "logs": logs[:100],
                     "tasks": weekly_tasks[:100], "project_items": weekly_items[:100]}
            prompt = "请将以下事实整理成工作周报。项目顺序必须与 projects 一致；没有证据的进度写待补充。\n" + json.dumps(facts, ensure_ascii=False, default=str)[:22000]
            try:
                candidate = call_model(profile, prompt, WEEKLY_SYSTEM_PROMPT)
                if valid_weekly_report(candidate, project_names):
                    body = candidate.rstrip() + "\n"
            except (OSError, ValueError, RuntimeError, urllib.error.URLError, TimeoutError):
                pass
        rid = uid()
        with db() as con:
            con.execute("INSERT INTO reports VALUES(?,?,?,?,?,?)", (rid, kind, start, end, body, now()))
        return one("SELECT * FROM reports WHERE id=?", (rid,))
    projects = sorted(set((x["project_name"] or "未归属") for x in logs + goals + tasks + items + file_events))
    lines = [f"# {kind}｜{start} 至 {(date.fromisoformat(end)-timedelta(days=1)).isoformat()}", "", f"> 生成于 {now()}。请核对后用于对外汇报。", ""]
    if date.fromisoformat(end) > date.today():
        lines += ["> 当前周期尚未结束，本报告仅截至生成时点。", ""]
    lines += ["## 本期目标", ""]
    lines += [f"- [{goal['status']}] {goal['title']}（{goal['project_name'] or '未归属'}）" for goal in goals] or ["本期尚未登记目标。"]
    lines.append("")
    lines += ["## 重点工作进展", ""]
    if not projects:
        lines.append("本周期暂无已记录的项目工作。")
    for name in projects:
        lines += [f"### {name}", ""]
        matching_logs = [x for x in logs if (x["project_name"] or "未归属") == name]
        matching_tasks = [x for x in tasks if (x["project_name"] or "未归属") == name]
        matching_items = [x for x in items if x["project_name"] == name]
        matching_files = [x for x in file_events if x["project_name"] == name]
        progress_start = len(lines)
        for item in matching_logs:
            lines.append(f"- {item['event_date']}：{item['body']}（工作记录 {item['id'][:8]}）")
        for item in matching_tasks:
            if item["completed_at"] and start <= item["completed_at"][:10] < end:
                lines.append(f"- 已完成事项：{item['title']}（事项 {item['id'][:8]}）")
        for item in matching_items:
            if not (start <= item["updated_at"][:10] < end):
                continue
            detail = f"；{item['outcome']}" if item["outcome"] else ""
            owner = f"；跟进 {item['owner']}" if item["owner"] else ""
            lines.append(f"- {item['kind']} [{item['status']}]：{item['title']}{detail}{owner}（台账 {item['id'][:8]}）")
        for item in matching_files:
            change = json.loads(item["payload"])
            for entry in change.get("changes", [])[:20]:
                lines.append(f"- 项目文件{entry['kind']}：`{entry['path']}`（仅表明文件变化，具体工作成果请核对）")
            if change.get("total", 0) > 20:
                lines.append(f"- 同次扫描另有 {change['total'] - 20} 项文件变化，详见项目档案。")
        if len(lines) == progress_start:
            lines.append("- 本周期暂无新增进展；遗留事项见下方。")
        lines.append("")
    open_tasks = [x for x in tasks if x["status"] not in ("已完成", "已取消")]
    open_items = [x for x in items if x["status"] not in ("已完成", "已关闭", "已取消")]
    lines += ["## 仍需推进与风险", ""]
    pending = [f"- [{x['status']}] {x['title']}" + (f"，跟进 {x['followup_date']}" if x["followup_date"] else "") for x in open_tasks]
    pending += [f"- {x['kind']} [{x['status']}] {x['title']}" + (f"，负责人/跟进人 {x['owner']}" if x["owner"] else "") for x in open_items]
    lines += pending or ["暂无已登记的待推进事项。"]
    lines += ["", "## 下期计划", "", "请结合最新情况填写下期交付结果及需要协调的事项。", ""]
    body = "\n".join(lines)
    profile = one("SELECT * FROM model_profiles WHERE id=?", (setting("active_model"),))
    if profile:
        eligible_logs = [x for x in logs if not x["local_only"] or profile["provider"] == "ollama"]
        eligible_tasks = [x for x in tasks if not x["local_only"] or profile["provider"] == "ollama"]
        eligible_items = [x for x in items if not x["local_only"] or profile["provider"] == "ollama"]
        eligible_goals = [x for x in goals if not x["local_only"] or profile["provider"] == "ollama"]
        context = json.dumps({"goals": eligible_goals[:50], "logs": eligible_logs[:100], "tasks": eligible_tasks[:100], "project_items": eligible_items[:100]}, ensure_ascii=False, default=str)[:22000]
        summary, note = model_summary("请根据以下事实写一段不超过300字的本期概况；区分完成与进行中；不得发明结果或数字；仅使用提供的资料。\n" + context)
        if summary:
            body = body.replace("## 重点工作进展", f"## AI 概况建议（请核对）\n\n{summary}\n\n## 重点工作进展", 1)
        if profile["provider"] != "ollama" and any(x["local_only"] for x in goals + logs + tasks + items):
            body = body.replace("## 重点工作进展", "> 仅本地项目未发送给云端模型，以下事实底稿仍含本地记录。\n\n## 重点工作进展", 1)
    rid = uid()
    with db() as con:
        con.execute("INSERT INTO reports VALUES(?,?,?,?,?,?)", (rid, kind, start, end, body, now()))
    return one("SELECT * FROM reports WHERE id=?", (rid,))


def report_docx(report: dict) -> bytes:
    """Export the saved, possibly edited report without generating new claims."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.4)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for line in report["body"].splitlines():
        if not line.strip():
            continue
        if line.startswith("### "):
            paragraph = doc.add_heading(line[4:], level=2)
        elif line.startswith("## "):
            paragraph = doc.add_heading(line[3:], level=1)
        elif line.startswith("# "):
            paragraph = doc.add_heading(line[2:], level=0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(line[2:], style="List Bullet")
        else:
            paragraph = doc.add_paragraph(line.removeprefix("> "))
        paragraph.paragraph_format.space_after = Pt(6)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


class Blob(ctypes.Structure):
    _fields_ = [("cbData", ctypes.c_ulong), ("pbData", ctypes.POINTER(ctypes.c_ubyte))]


def protect_key(secret: str) -> str:
    if not secret:
        return ""
    if os.name != "nt":
        raise RuntimeError("当前环境无法安全保存云端密钥")
    raw = secret.encode("utf-8")
    inp_buffer = ctypes.create_string_buffer(raw)
    inp = Blob(len(raw), ctypes.cast(inp_buffer, ctypes.POINTER(ctypes.c_ubyte)))
    out = Blob()
    if not ctypes.windll.crypt32.CryptProtectData(ctypes.byref(inp), None, None, None, None, 0, ctypes.byref(out)):
        raise RuntimeError("Windows 凭据加密失败")
    try:
        return base64.b64encode(ctypes.string_at(out.pbData, out.cbData)).decode()
    finally:
        ctypes.windll.kernel32.LocalFree(out.pbData)


def unprotect_key(blob: str) -> str:
    if not blob:
        return ""
    raw = base64.b64decode(blob)
    inp_buffer = ctypes.create_string_buffer(raw)
    inp = Blob(len(raw), ctypes.cast(inp_buffer, ctypes.POINTER(ctypes.c_ubyte)))
    out = Blob()
    if not ctypes.windll.crypt32.CryptUnprotectData(ctypes.byref(inp), None, None, None, None, 0, ctypes.byref(out)):
        raise RuntimeError("无法解锁 API Key")
    try:
        return ctypes.string_at(out.pbData, out.cbData).decode("utf-8")
    finally:
        ctypes.windll.kernel32.LocalFree(out.pbData)


def model_public(profile):
    return {k: profile[k] for k in ("id", "name", "provider", "base_url", "model", "created_at")}


def create_profile(data: dict):
    provider = data.get("provider")
    if provider not in ("ollama", "online"):
        raise ValueError("请选择本地或云端模型")
    base = (data.get("base_url") or ("http://127.0.0.1:11434" if provider == "ollama" else "")).strip().rstrip("/")
    if provider == "ollama":
        from urllib.parse import urlparse
        parts = urlparse(base)
        if parts.hostname not in ("localhost", "127.0.0.1", "::1"):
            raise ValueError("本地 Ollama 只允许本机地址")
    elif not base.startswith("https://"):
        raise ValueError("云端 API 地址必须使用 HTTPS")
    name = str(data.get("name") or "").strip()
    model = str(data.get("model") or "").strip()
    if not name or not model:
        raise ValueError("请输入连接名称与模型名称")
    secret = str(data.get("api_key") or "")
    if provider == "online" and not secret:
        raise ValueError("请输入云端 API Key")
    pid = uid()
    with db() as con:
        con.execute("INSERT INTO model_profiles VALUES(?,?,?,?,?,?,?)", (pid, name, provider, base, model, protect_key(secret), now()))
    return model_public(one("SELECT * FROM model_profiles WHERE id=?", (pid,)))


def call_model(profile: dict, prompt: str, system_prompt: str = "") -> str:
    provider = profile["provider"]
    messages = ([{"role": "system", "content": system_prompt}] if system_prompt else []) + [{"role": "user", "content": prompt}]
    if provider == "ollama":
        from urllib.parse import urlparse
        if urlparse(profile["base_url"]).hostname not in ("localhost", "127.0.0.1", "::1"):
            raise ValueError("Ollama 地址不在本机")
        if "cloud" in profile["model"].lower():
            raise ValueError("严格本地模式不允许选择云端模型名称")
        if profile["model"] not in available_ollama(profile["base_url"]):
            raise ValueError("该模型不在本机 Ollama 已安装列表中")
        endpoint = profile["base_url"] + "/api/chat"
        request_body = {"model": profile["model"], "messages": messages, "stream": False}
        headers = {"Content-Type": "application/json"}
    else:
        if not profile["base_url"].startswith("https://"):
            raise ValueError("云端模型需使用 HTTPS")
        endpoint = profile["base_url"] + "/chat/completions"
        request_body = {"model": profile["model"], "messages": messages, "stream": False}
        headers = {"Content-Type": "application/json", "Authorization": "Bearer " + unprotect_key(profile["key_blob"])}
    req = urllib.request.Request(endpoint, data=json.dumps(request_body, ensure_ascii=False).encode("utf-8"), headers=headers, method="POST")
    class NoRedirect(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, request, fp, code, msg, headers, newurl):
            raise ValueError("模型地址发生跳转，请核对 API 地址")

    with urllib.request.build_opener(NoRedirect).open(req, timeout=80) as response:
        answer = json.loads(response.read(4 * 1024 * 1024))
    if provider == "ollama":
        return str(answer.get("message", {}).get("content") or "").strip()
    return str(answer.get("choices", [{}])[0].get("message", {}).get("content") or "").strip()


def model_list():
    return [model_public(p) for p in rows("SELECT * FROM model_profiles ORDER BY created_at DESC")]


def available_ollama(base="http://127.0.0.1:11434"):
    from urllib.parse import urlparse
    if urlparse(base).hostname not in ("localhost", "127.0.0.1", "::1"):
        raise ValueError("只允许检测本机 Ollama")
    try:
        with urllib.request.urlopen(base.rstrip("/") + "/api/tags", timeout=5) as response:
            data = json.loads(response.read(2 * 1024 * 1024))
    except (OSError, ValueError) as exc:
        raise ValueError("未连接到 Ollama。请确认它已启动，并检查本机地址与端口。") from exc
    return [m.get("name") for m in data.get("models", []) if m.get("name") and "cloud" not in m.get("name", "").lower()]


def backup_data():
    backups = DATA / "backups"
    backups.mkdir(exist_ok=True)
    target = backups / f"workbench-{datetime.now():%Y%m%d-%H%M%S}.sqlite3"
    with LOCK, sqlite3.connect(DB_PATH) as source, sqlite3.connect(target) as dest:
        source.backup(dest)
    return {"path": str(target), "note": "仅备份工作台数据库；项目原文件请另行备份"}
